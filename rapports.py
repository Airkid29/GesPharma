import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

def exporter_ventes_excel(liste_ventes, nom_fichier="ventes_du_jour.xlsx"):
    """
    Exporte la liste des ventes vers un fichier Excel avec mise en forme.

    Args:
        liste_ventes (list): La liste des dictionnaires de ventes.
        nom_fichier (str, optional): Le nom du fichier Excel à créer.
                                    Par défaut, "ventes_du_jour.xlsx".
    """
    if not liste_ventes:
        print("Aucune vente à exporter.")
        return

    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.title = "Ventes"

    # Définition des styles
    style_titre = Font(name='Arial', size=16, bold=True, color='FFFFFF')
    fill_titre = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid')
    
    style_entete = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    fill_entete = PatternFill(start_color='95B3D7', end_color='95B3D7', fill_type='solid')
    align_center = Alignment(horizontal='center', vertical='center')
    
    border_thin = Border(left=Side(style='thin'), 
                         right=Side(style='thin'), 
                         top=Side(style='thin'), 
                         bottom=Side(style='thin'))

    # 1. Titre du rapport
    feuille.merge_cells('A1:E1')
    cell_titre = feuille['A1']
    cell_titre.value = f"Rapport de Ventes - {datetime.now().strftime('%d/%m/%Y')}"
    cell_titre.font = style_titre
    cell_titre.fill = fill_titre
    cell_titre.alignment = align_center

    # 2. En-têtes de colonnes
    en_tetes = ["Date", "Médicament", "Quantité", "Prix Unitaire", "Prix Total"]
    for col_num, en_tete in enumerate(en_tetes, 1):
        cell = feuille.cell(row=2, column=col_num)
        cell.value = en_tete
        cell.font = style_entete
        cell.fill = fill_entete
        cell.alignment = align_center
        cell.border = border_thin

    # 3. Données
    total_general = 0.0
    row_num = 3
    for vente in liste_ventes:
        # Date
        cell_date = feuille.cell(row=row_num, column=1)
        cell_date.value = vente.get('date', '')
        cell_date.border = border_thin
        
        # Médicament
        cell_nom = feuille.cell(row=row_num, column=2)
        cell_nom.value = vente.get('medicament', '')
        cell_nom.border = border_thin
        
        # Quantité
        cell_qte = feuille.cell(row=row_num, column=3)
        cell_qte.value = vente.get('quantite', 0)
        cell_qte.alignment = Alignment(horizontal='center')
        cell_qte.border = border_thin
        
        # Prix Unit
        cell_pu = feuille.cell(row=row_num, column=4)
        cell_pu.value = float(vente.get('prix_unitaire', 0))
        cell_pu.number_format = '#,##0.00 "€"' 
        cell_pu.border = border_thin
        
        # Prix Total
        p_total = float(vente.get('prix_total', 0))
        cell_pt = feuille.cell(row=row_num, column=5)
        cell_pt.value = p_total
        cell_pt.number_format = '#,##0.00 "€"'
        cell_pt.font = Font(bold=True)
        cell_pt.border = border_thin
        
        total_general += p_total
        row_num += 1

    # 4. Ligne Total Général
    row_total = row_num
    feuille.merge_cells(f'A{row_total}:D{row_total}')
    cell_label_total = feuille[f'A{row_total}']
    cell_label_total.value = "TOTAL GÉNÉRAL"
    cell_label_total.font = Font(bold=True)
    cell_label_total.alignment = Alignment(horizontal='right')
    cell_label_total.border = border_thin
    
    # Bordure pour la zone fusionnée
    for col in range(1, 5):
        cell = feuille.cell(row=row_total, column=col)
        cell.border = border_thin

    cell_val_total = feuille.cell(row=row_total, column=5)
    cell_val_total.value = total_general
    cell_val_total.font = Font(bold=True, size=12, color='FF0000') 
    cell_val_total.number_format = '#,##0.00 "€"'
    cell_val_total.border = border_thin

    # 5. Ajustement largeurs
    for column_cells in feuille.columns:
        length = max(len(str(cell.value)) if cell.value else 0 for cell in column_cells)
        feuille.column_dimensions[get_column_letter(column_cells[0].column)].width = length + 4

    try:
        classeur.save(nom_fichier)
        print(f"Les ventes ont été exportées avec succès vers '{nom_fichier}'.")
    except Exception as e:
        print(f"Erreur lors de l'exportation vers Excel : {e}")

if __name__ == '__main__':
    # Données de test pour vérification immédiate
    liste_test = [
        {"date": "2023-10-27 10:00", "medicament": "Doliprane 1000mg", "quantite": 2, "prix_unitaire": 2.50, "prix_total": 5.00},
        {"date": "2023-10-27 10:05", "medicament": "Spasfon", "quantite": 1, "prix_unitaire": 4.50, "prix_total": 4.50},
        {"date": "2023-10-27 11:30", "medicament": "Advil 400mg", "quantite": 3, "prix_unitaire": 3.90, "prix_total": 11.70},
    ]
    exporter_ventes_excel(liste_test, "test_ventes_style.xlsx")